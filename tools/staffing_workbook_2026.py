#!/usr/bin/env python3
"""Extract the Sept 2026 PPS staffing / enrollment workbook ("Board Request as of 9/22").

  python3 tools/staffing_workbook_2026.py

Reads  sources/staffing/2026-09-22_board-request/board_request_as_of_2026-09-22.xlsx
Writes sources/staffing/2026-09-22_board-request/tabs/NN_<tab>.csv   (mechanical cell dump of every tab, cached values)
       analysis/data/enrollment_2026-27_sept_by_grade.csv         (the "26-27 As of 9/23 Counts" tab, one row per school/program/grade)
       analysis/data/enrollment_2026-27_sept_by_school.csv        (per-school totals joined to staffing counts, Oct 2025 counts
                                                                    and the family-guide status-quo projection)
Citations: [PPS:staffing/2026-09-22_board-request/board_request_as_of_2026-09-22.xlsx#<tab>!R<row>]
"""
import csv, re, datetime, warnings
from collections import defaultdict
from pathlib import Path
import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "sources/staffing/2026-09-22_board-request"
XLSX = SRC / "board_request_as_of_2026-09-22.xlsx"
REL = "staffing/2026-09-22_board-request/board_request_as_of_2026-09-22.xlsx"
COUNTS, DETAIL = "26-27 As of 923 Counts", "School Detail wo ESL or SPED or"

def cite(tab, row): return f"[PPS:{REL}#{tab}!R{row}]"

def cell(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return "" if v is None else str(v)

def span(v):
    """Excel turned grade spans like "6-8" into dates (2026-06-08); undo that for the derived tables only."""
    if isinstance(v, datetime.datetime): return f"{v.month}-{v.day}"
    return cell(v)

# workbook name -> (slug of analysis/schools page, name used in analysis/data/school_metrics.csv)
ALIAS = {"Access": ("access", "ACCESS Academy"), "Boise-Eliot/Humboldt": ("boise-eliot-humboldt", None),
         "César Chávez": ("cesar-chavez", None), "MLK Jr": ("dr-martin-luther-king-jr", "Dr. Martin Luther King Jr."),
         "Metro. Learning Center": ("metropolitan-learning-center", "Metropolitan Learning Center (MLC)"),
         "Sunnyside Environmental": ("sunnyside", None), "Lee": ("sunrise", 'Lee (Jason Lee; renamed - "Sunrise")'),
         "Sunrise": ("sunrise", 'Lee (Jason Lee; renamed - "Sunrise")'), "Lane": ("brentwood", "Lane MS (renamed Brentwood MS)"),
         "Mt Tabor": ("mt-tabor", "Mt. Tabor MS"), "da Vinci": ("da-vinci", "da Vinci Arts MS"), "Gray": ("gray", "Robert Gray MS"),
         "Ida B. Wells-Barnett": ("ida-b-wells", "Ida B. Wells HS"), "Benson": ("benson", "Benson Polytechnic HS")}
MS = {"Beaumont", "George", "Harriet Tubman", "Harrison Park", "Hosford", "Jackson", "Kellogg", "Ockley Green", "Roseway Heights",
      "Sellwood", "West Sylvan"}
HS = {"Cleveland", "Franklin", "Grant", "Jefferson", "Lincoln", "McDaniel", "Roosevelt"}

def slug_and_metric_name(name):
    if name in ALIAS:
        s, m = ALIAS[name]; return s, (m or name)
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return s, (f"{name} MS" if name in MS else f"{name} HS" if name in HS else name)

def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    out = SRC / "tabs"; out.mkdir(exist_ok=True)
    for i, ws in enumerate(wb.worksheets):
        fn = out / f"{i:02d}_{re.sub(r'[^A-Za-z0-9]+', '_', ws.title).strip('_')}.csv"
        with fn.open("w", newline="") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow([cell(v) for v in row])

    # --- per grade (counts tab)
    ws = wb[COUNTS]
    hdr = [c.value for c in ws[1]]
    assert hdr[:9] == ["schoolname", "subgroupcd", "grade", "unit", "blend applied", "initial sections", "Enrollment Count",
                       "school Funded Homerooms", "Projected Enrollment"], hdr
    grade_rows = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row[0]: continue
        slug, _ = slug_and_metric_name(row[0])
        grade_rows.append({"school": row[0], "slug": slug, "subgroup": row[1], "grade": cell(row[2]), "unit": cell(row[3]),
                           "blend_applied": row[4], "initial_sections": cell(row[5]),
                           "enrollment_count_2026_09": cell(row[6]), "school_funded_homerooms": cell(row[7]),
                           "projected_enrollment_2026_27": cell(row[8]), "source": cite(COUNTS, r)})
    with (ROOT / "analysis/data/enrollment_2026-27_sept_by_grade.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(grade_rows[0])); w.writeheader(); w.writerows(grade_rows)

    agg = defaultdict(lambda: {"k5": 0, "g68": 0, "proj": 0, "blend": 0, "rows": [], "grades": set(), "subgroups": set()})
    for g in grade_rows:
        a = agg[g["school"]]; n = int(float(g["enrollment_count_2026_09"] or 0))
        a["k5" if g["grade"] in ("KG", "1", "2", "3", "4", "5") else "g68"] += n
        a["proj"] += int(float(g["projected_enrollment_2026_27"] or 0))
        a["blend"] += g["blend_applied"] == "Yes"; a["grades"].add(g["grade"]); a["subgroups"].add(g["subgroup"])
        a["rows"].append(int(g["source"].rsplit("!R", 1)[1][:-1]))

    # --- staffing counts for every school (School Detail tab)
    det = {}
    ws = wb[DETAIL]
    for r in range(5, ws.max_row + 1):
        name = ws.cell(r, 5).value
        if not name or "TOTAL" in str(name).upper() or "Other MPG" in str(name): continue
        det[name] = {"row": r, "grade_span": span(ws.cell(r, 4).value), "title_I": cell(ws.cell(r, 1).value),
                     "ode_id": cell(ws.cell(r, 3).value), "staff_2025_26": cell(ws.cell(r, 11).value),
                     "staff_2026_27": cell(ws.cell(r, 17).value)}

    metrics = {(m["school"]): m for m in csv.DictReader(open(ROOT / "analysis/data/school_metrics.csv"))
               if m["metric"] == "enrollment" and m["as_of"].startswith("2025-10-01")}
    impact = {(x["slug"], x["scenario"]): x for x in csv.DictReader(open(ROOT / "analysis/data/school_impact.csv"))}

    names = list(det) + [n for n in agg if n not in det and not (n == "Sunrise" and "Lee" in det)]
    rows = []
    for name in names:
        d = det.get(name) or {}; a = agg.get(name) or agg.get("Sunrise" if name == "Lee" else name)
        slug, mname = slug_and_metric_name(name)
        m = metrics.get(mname); sq = impact.get((slug, "StatusQuo")); A = impact.get((slug, "A")); B = impact.get((slug, "B"))
        sept = (a["k5"] + a["g68"]) if a else None
        cites = []
        if a: cites.append(f"[PPS:{REL}#{COUNTS}!R{min(a['rows'])}-R{max(a['rows'])}]")
        if d: cites.append(cite(DETAIL, d["row"]))
        comparable = bool(a and m and d.get("grade_span") not in ("K-12",))
        rows.append({
            "school": name + (" (shown as Sunrise in the counts tab)" if name == "Lee" else ""), "slug": slug,
            "grade_span": d.get("grade_span", ""), "title_I_2026_27": d.get("title_I", ""),
            "sept_2026_enrollment": sept if sept is not None else "",
            "sept_2026_k5": a["k5"] if a else "", "sept_2026_6_8": a["g68"] if a and a["g68"] else "",
            "spring_projection_2026_27": a["proj"] if a else "",
            "sept_minus_projection": (sept - a["proj"]) if a else "",
            "blended_grade_rows": a["blend"] if a else "",
            "staffing_count_2025_26": d.get("staff_2025_26", ""), "staffing_count_2026_27": d.get("staff_2026_27", ""),
            "enrollment_2025_10_01": m["value"] if m else "",
            "change_vs_2025_10_01": (sept - int(m["value"])) if comparable else "",
            "guide_status_quo_2027_28": sq["proj_enr_2027_28"] if sq else "",
            "scenario_A": A["proposed_action"] if A else "", "scenario_B": B["proposed_action"] if B else "",
            "note": ("K-8 grades only; Oct 2025 figure includes grades 9-12, so no change is computed" if d.get("grade_span") == "K-12" else
                     "No Sept 2026 actuals in the workbook (middle and high schools only have staffing counts)" if not a else ""),
            "source": " ".join(cites + ([m["source"]] if m else []) +
                               ([f"[PPS:pps_net/rightsizing/school_impact/{slug.replace('-', '_')}.pdf p.4]"] if sq else [])),
        })
    with (ROOT / "analysis/data/enrollment_2026-27_sept_by_school.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    print(f"{len(wb.worksheets)} tabs dumped; {len(grade_rows)} grade rows; {len(rows)} schools")

if __name__ == "__main__":
    main()
